import os
import re
import fitz  # PyMuPDF
from datetime import datetime
from typing import List, Literal, Optional
from pydantic import BaseModel, Field
from google import genai
from google.genai import types
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# =====================================================================
# 1. ESQUEMA DE VALIDAÇÃO ESTRITA (PYDANTIC)
# Bloqueia unidades falsas (ex: "Sesc São Bento") e textos truncados
# =====================================================================

UnidadesCapital = Literal[
    "Sesc 24 de Maio",
    "Sesc Avenida Paulista",
    "Sesc Belenzinho",
    "Sesc Bom Retiro",
    "Sesc Campo Limpo",
    "Sesc Carmo",
    "Sesc Casa Verde",
    "Sesc CineSesc",
    "Sesc Consolação",
    "Sesc Florêncio de Abreu",
    "Sesc Interlagos",
    "Sesc Ipiranga",
    "Sesc Itaquera",
    "Sesc Pinheiros",
    "Sesc Pompeia",
    "Sesc Santana",
    "Sesc Santo Amaro",
    "Sesc Vila Mariana",
]

class EventoMusical(BaseModel):
    artista_titulo: str = Field(
        description=(
            "APENAS o nome do artista, banda ou título oficial do show (em destaque no topo da nota). "
            "NUNCA inclua frases da sinopse (ex: 'Grupo apresenta...', 'Show inspirado...'), "
            "nem duração ('60 min.'), nem códigos como 'AL', 'VIVÊNCIA' ou 'BATE-PAPO'."
        )
    )
    unidade: UnidadesCapital = Field(
        description="Unidade oficial do Sesc na Capital. Ignore estações de metrô como 'Metrô São Bento'."
    )
    data: str = Field(
        description="Datas do show no formato DD/MM (ex: '24/10', '9 e 10/10', '1 a 12/10')."
    )
    horario: Optional[str] = Field(
        default=None,
        description="Horário do show (ex: '20h', '21h', '19h30'). Deixe null se não informado."
    )
    preco: Optional[str] = Field(
        default=None,
        description="Preços no formato 'R$XX / R$XX / R$XX' ou 'GRÁTIS'."
    )
    classificacao: Optional[Literal["AL", "A10", "A12", "A14", "A16", "A18"]] = Field(
        default=None,
        description="Classificação indicativa oficial: AL, A10, A12, A14, A16 ou A18."
    )

class ListaProgramacao(BaseModel):
    eventos: List[EventoMusical]


# =====================================================================
# 2. FILTRO DE PÁGINAS COM PYMUPDF
# Descarta páginas de Teatro, Dança, Turismo e Infantil antes da IA ler
# =====================================================================

def filtrar_paginas_musica(caminho_pdf_entrada: str, caminho_pdf_saida: str) -> int:
    doc = fitz.open(caminho_pdf_entrada)
    novo_doc = fitz.open()

    editorias_excluir = [
        "TEATRO", "DANÇA", "CIRCO", "CINEMA", "ARTES VISUAIS",
        "LITERATURA", "CRIANÇAS", "TURISMO", "IDOSOS", "ESPORTE", "JOVENS"
    ]

    em_secao_musica = False
    paginas_musica = []

    for i in range(len(doc)):
        pagina = doc[i]
        # Lê os blocos superiores e laterais onde fica a tarja da editoria na revista
        blocos = pagina.get_text("blocks")
        texto_cabecalho = " ".join([b[4].upper() for b in blocos if b[1] < 130 or b[0] < 70 or b[0] > pagina.rect.width - 70])

        if re.search(r"\bM[ÚU]SICA\b", texto_cabecalho):
            em_secao_musica = True
        elif em_secao_musica and any(re.search(rf"\b{ed}\b", texto_cabecalho) for ed in editorias_excluir):
            em_secao_musica = False

        if em_secao_musica:
            paginas_musica.append(i)

    # Caso o PDF daquele mês tenha as tarjas de seção convertidas em imagem/curva,
    # envia o PDF completo para não perder nenhum evento
    if not paginas_musica:
        print("Tarja textual 'MÚSICA' não detectada isoladamente; enviando PDF completo para a IA filtrar.")
        doc.save(caminho_pdf_saida)
        total = len(doc)
        doc.close()
        return total

    for p in paginas_musica:
        novo_doc.insert_pdf(doc, from_page=p, to_page=p)

    novo_doc.save(caminho_pdf_saida)
    total = len(novo_doc)
    novo_doc.close()
    doc.close()
    return total


# =====================================================================
# 3. EXTRAÇÃO ESTRUTURADA COM GEMINI 2.5 FLASH
# =====================================================================

def extrair_eventos_pdf(caminho_pdf: str) -> List[EventoMusical]:
    api_key = os.environ.get("GEMINI_API_KEY")
    if not api_key:
        raise ValueError("A variável de ambiente GEMINI_API_KEY não foi configurada nos Secrets do GitHub!")

    client = genai.Client(api_key=api_key)
    pdf_filtrado = "temp_secao_musica.pdf"

    qtd_paginas = filtrar_paginas_musica(caminho_pdf, pdf_filtrado)
    print(f"Enviando {qtd_paginas} página(s) para extração estruturada com Gemini 2.5 Flash...")

    arquivo_remoto = client.files.upload(file=pdf_filtrado)

    prompt = """
    Você está lendo a revista 'Em Cartaz' do Sesc São Paulo.
    Extraia TODOS os shows da seção de MÚSICA que ocorrem APENAS nas 18 unidades da CAPITAL de São Paulo.

    REGRAS OBRIGATÓRIAS:
    1. Cada página possui 3 colunas verticais. Leia cada coluna verticalmente de cima para baixo, sem jamais juntar textos de colunas vizinhas.
    2. No campo 'artista_titulo', extraia SOMENTE o título do show ou nome do artista/banda (que fica em negrito/destaque no início do bloco).
       PROIBIDO incluir frases da descrição/sinopse (ex: 'Grupo apresenta...', 'Show inspirado no...', 'Reúne músicas...').
       PROIBIDO incluir horários, preços, durações ('60 min.') ou classificações ('AL', 'A12') dentro de 'artista_titulo'.
    3. EXCLUA qualquer evento de Teatro, Dança, Circo, Infantil/Crianças, Contação de Histórias, Bate-papo, Cursos, Oficinas, Vivências ou Turismo Social.
    4. EXCLUA unidades do interior, litoral ou Grande SP (como Santo André, São Caetano, Guarulhos, Osasco, Mogi das Cruzes, Jundiaí, etc.).
    5. Atenção: 'Metrô São Bento', 'Metrô República' ou 'Casa Verde' no meio de um endereço são referências de localização; classifique apenas pela tarja oficial da Unidade do Sesc.
    """

    try:
        resposta = client.models.generate_content(
            model="gemini-2.5-flash",
            contents=[arquivo_remoto, prompt],
            config=types.GenerateContentConfig(
                temperature=0.0,
                response_mime_type="application/json",
                response_schema=ListaProgramacao,
            ),
        )
        resultado = ListaProgramacao.model_validate_json(resposta.text)
    finally:
        # Limpa os arquivos temporários na nuvem e localmente
        client.files.delete(name=arquivo_remoto.name)
        if os.path.exists(pdf_filtrado):
            os.remove(pdf_filtrado)

    # Remove eventuais duplicatas e ordena por Unidade e Data
    vistos = set()
    eventos_unicos = []
    for ev in resultado.eventos:
        chave = (ev.unidade, ev.artista_titulo.lower().strip(), ev.data.strip())
        if chave not in vistos:
            vistos.add(chave)
            eventos_unicos.append(ev)

    return sorted(eventos_unicos, key=lambda x: (x.unidade, x.data))


# =====================================================================
# 4. GERAÇÃO DO EXCEL NO PADRÃO VISUAL DO REPOSITÓRIO
# =====================================================================

def salvar_planilha_excel(eventos: List[EventoMusical], mes_ano_label: str, caminho_xlsx: str):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Programação Musical"

    # Linha 1: Título Principal
    ws.merge_cells("A1:F1")
    ws["A1"] = f"Programação Musical – Sesc SP Capital | {mes_ano_label}"
    ws["A1"].font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="C00000")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    # Linha 2: Metadados da Extração
    ws.merge_cells("A2:F2")
    data_extracao = datetime.now().strftime("%d/%m/%Y %H:%M")
    ws["A2"] = f"Extraído em: {data_extracao}   |   Fonte: PDF Em Cartaz – sescsp.org.br"
    ws["A2"].font = Font(name="Calibri", size=10, italic=True, color="555555")
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 20

    # Linha 3: Cabeçalhos da Tabela
    colunas = ["Artista / Título", "Unidade", "Data", "Horário", "Preço", "Classificação"]
    for idx_col, nome_col in enumerate(colunas, start=1):
        celula = ws.cell(row=3, column=idx_col, value=nome_col)
        celula.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        celula.fill = PatternFill("solid", fgColor="8B0000")
        celula.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[3].height = 24

    # Linhas de Dados
    borda = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    for idx_linha, ev in enumerate(eventos, start=4):
        cor_zebra = "F5F5F5" if idx_linha % 2 == 0 else "FFFFFF"
        preenchimento = PatternFill("solid", fgColor=cor_zebra)
        valores = [ev.artista_titulo, ev.unidade, ev.data, ev.horario, ev.preco, ev.classificacao]

        for idx_col, valor in enumerate(valores, start=1):
            c = ws.cell(row=idx_linha, column=idx_col, value=valor)
            c.number_format = "@"  # Trava formato como texto puro (evita que '9/10' vire data errada)
            c.fill = preenchimento
            c.border = borda
            c.alignment = Alignment(
                horizontal="left" if idx_col in (1, 2) else "center",
                vertical="center",
                wrap_text=True,
            )

    # Rodapé com Totalizador
    linha_final = len(eventos) + 4
    ws.merge_cells(f"A{linha_final}:F{linha_final}")
    ws[f"A{linha_final}"] = f"Total de eventos musicais: {len(eventos)}"
    ws[f"A{linha_final}"].font = Font(name="Calibri", size=10, bold=True)
    ws[f"A{linha_final}"].alignment = Alignment(horizontal="right", vertical="center")

    # Ajuste das Larguras das Colunas
    larguras = {"A": 42.0, "B": 24.0, "C": 22.0, "D": 14.0, "E": 22.0, "F": 16.0}
    for col_letra, largura in larguras.items():
        ws.column_dimensions[col_letra].width = largura

    wb.save(caminho_xlsx)
    print(f"Planilha gerada com sucesso: {caminho_xlsx} ({len(eventos)} shows válidos)")
